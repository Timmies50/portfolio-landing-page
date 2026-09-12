from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).parent
OUTPUT = ROOT / "excel" / "sql-projects-report.xlsx"
OUTPUT.parent.mkdir(parents=True, exist_ok=True)

NAVY = "1F3A93"
TEAL = "00D9FF"
LIGHT_BLUE = "EAF3FF"
PALE_TEAL = "E8FBFF"
SLATE = "334155"
GREEN = "DCFCE7"
AMBER = "FEF3C7"
WHITE = "FFFFFF"
GRID = "CBD5E1"

thin_gray = Side(style="thin", color=GRID)
medium_navy = Side(style="medium", color=NAVY)

wb = Workbook()
ws_inputs = wb.active
ws_inputs.title = "Inputs"
ws_calcs = wb.create_sheet("Calcs")
ws_outputs = wb.create_sheet("Outputs")

for ws in (ws_inputs, ws_calcs, ws_outputs):
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A4"

# Inputs sheet
ws_inputs["A1"] = "SQL Projects Workbook"
ws_inputs["A1"].font = Font(size=18, bold=True, color=WHITE)
ws_inputs["A1"].fill = PatternFill("solid", fgColor=NAVY)
ws_inputs.merge_cells("A1:F1")
ws_inputs["A2"] = "Inputs"
ws_inputs["A2"].font = Font(size=12, bold=True, color=NAVY)
ws_inputs["B2"] = 0
ws_inputs["B2"].number_format = '#,##0'
ws_inputs["C2"] = "Minimum order amount scenario"
ws_inputs["C2"].font = Font(italic=True, color=SLATE)
ws_inputs["A3"] = "Change the value in B2 to show only orders at or above that amount."
ws_inputs["A3"].font = Font(italic=True, color=SLATE)
ws_inputs.merge_cells("A3:F3")

ws_inputs["A5"] = "Customer ID"
ws_inputs["B5"] = "Customer Name"
ws_inputs["C5"] = "City"
customers = [
    (101, "Ada", "Lagos"),
    (102, "John", "Abuja"),
    (103, "Mercy", "Owerri"),
    (104, "David", "Enugu"),
    (105, "Grace", "Lagos"),
]
for row_num, row in enumerate(customers, start=6):
    for col_num, value in enumerate(row, start=1):
        ws_inputs.cell(row_num, col_num, value)

ws_inputs["E5"] = "Order ID"
ws_inputs["F5"] = "Customer ID"
ws_inputs["G5"] = "Product"
ws_inputs["H5"] = "Total Amount"
orders = [
    (1001, 101, "Laptop", 450000),
    (1002, 102, "Phone", 180000),
    (1003, 101, "Monitor", 150000),
    (1004, 105, "Tablet", 120000),
    (1005, 103, "Laptop", 450000),
]
for row_num, row in enumerate(orders, start=6):
    for col_num, value in enumerate(row, start=5):
        ws_inputs.cell(row_num, col_num, value)
    ws_inputs.cell(row_num, 8).number_format = '#,##0'

for cell in ws_inputs[5]:
    if cell.value is not None:
        cell.font = Font(bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(horizontal="center")

customers_table = Table(displayName="Customers", ref="A5:C10")
customers_table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
ws_inputs.add_table(customers_table)
orders_table = Table(displayName="Orders", ref="E5:H10")
orders_table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
ws_inputs.add_table(orders_table)

threshold_validation = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=False)
threshold_validation.error = "Enter a number equal to or greater than zero."
threshold_validation.errorTitle = "Invalid scenario value"
threshold_validation.prompt = "Use 0 for all orders, or enter a higher minimum order amount."
threshold_validation.promptTitle = "Scenario control"
ws_inputs.add_data_validation(threshold_validation)
threshold_validation.add(ws_inputs["B2"])

# Calcs sheet
ws_calcs["A1"] = "Calculation Layer"
ws_calcs["A1"].font = Font(size=18, bold=True, color=WHITE)
ws_calcs["A1"].fill = PatternFill("solid", fgColor=NAVY)
ws_calcs.merge_cells("A1:J1")
headers = ["Order ID", "Customer ID", "Product", "Order Amount", "Customer Name", "City", "Order Key", "Scenario", "First Customer Order", "Display Order"]
for col_num, header in enumerate(headers, start=1):
    cell = ws_calcs.cell(3, col_num, header)
    cell.font = Font(bold=True, color=WHITE)
    cell.fill = PatternFill("solid", fgColor=SLATE)
    cell.alignment = Alignment(horizontal="center", wrap_text=True)

for row_num in range(4, 9):
    input_row = row_num + 2
    ws_calcs.cell(row_num, 1, f"=Inputs!E{input_row}")
    ws_calcs.cell(row_num, 2, f"=Inputs!F{input_row}")
    ws_calcs.cell(row_num, 3, f"=Inputs!G{input_row}")
    ws_calcs.cell(row_num, 4, f"=Inputs!H{input_row}")
    ws_calcs.cell(row_num, 5, f'=IFERROR(XLOOKUP(B{row_num},Inputs!$A$6:$A$10,Inputs!$B$6:$B$10,""),"")')
    ws_calcs.cell(row_num, 6, f'=IFERROR(XLOOKUP(B{row_num},Inputs!$A$6:$A$10,Inputs!$C$6:$C$10,""),"")')
    ws_calcs.cell(row_num, 7, f'=IFERROR(A{row_num}&"-"&B{row_num},"")')
    ws_calcs.cell(row_num, 8, f'=IF(D{row_num}>=Inputs!$B$2,"Include","Exclude")')
    ws_calcs.cell(row_num, 9, f'=IF(COUNTIFS($B$4:B{row_num},B{row_num})=1,"Yes","No")')
    ws_calcs.cell(row_num, 10, f'=IF(H{row_num}="Include",COUNTIFS($H$4:H{row_num},"Include"),"")')
    ws_calcs.cell(row_num, 4).number_format = '#,##0'

for col in range(1, 11):
    for row in range(4, 9):
        ws_calcs.cell(row, col).border = Border(bottom=thin_gray)

ws_calcs.conditional_formatting.add("H4:H8", CellIsRule(operator="equal", formula=['"Include"'], fill=PatternFill("solid", fgColor=GREEN)))
ws_calcs.conditional_formatting.add("H4:H8", CellIsRule(operator="equal", formula=['"Exclude"'], fill=PatternFill("solid", fgColor=AMBER)))

# Outputs sheet
ws_outputs["A1"] = "SQL Projects Report"
ws_outputs["A1"].font = Font(size=18, bold=True, color=WHITE)
ws_outputs["A1"].fill = PatternFill("solid", fgColor=NAVY)
ws_outputs.merge_cells("A1:E1")
ws_outputs["A2"] = "Customers with orders, filtered by the selected minimum order amount"
ws_outputs["A2"].font = Font(italic=True, color=SLATE)
ws_outputs.merge_cells("A2:E2")

kpis = [
    ("A4", "Orders shown", "B4", '=COUNTIFS(Calcs!$H$4:$H$8,"Include")', '#,##0'),
    ("C4", "Revenue shown", "D4", '=SUMIFS(Calcs!$D$4:$D$8,Calcs!$H$4:$H$8,"Include")', '#,##0'),
    ("A6", "Customers shown", "B6", '=COUNTIFS(Calcs!$I$4:$I$8,"Yes",Calcs!$H$4:$H$8,"Include")', '#,##0'),
    ("C6", "Average order", "D6", '=IFERROR(D4/B4,0)', '#,##0'),
]
for label_cell, label, value_cell, formula, number_format in kpis:
    ws_outputs[label_cell] = label
    ws_outputs[label_cell].font = Font(bold=True, color=WHITE)
    ws_outputs[label_cell].fill = PatternFill("solid", fgColor=SLATE)
    ws_outputs[value_cell] = formula
    ws_outputs[value_cell].font = Font(size=14, bold=True, color=NAVY)
    ws_outputs[value_cell].fill = PatternFill("solid", fgColor=PALE_TEAL)
    ws_outputs[value_cell].number_format = number_format
    ws_outputs[label_cell].alignment = Alignment(horizontal="center")
    ws_outputs[value_cell].alignment = Alignment(horizontal="center")

ws_outputs["A8"] = "Scenario minimum"
ws_outputs["A8"].font = Font(bold=True, color=WHITE)
ws_outputs["A8"].fill = PatternFill("solid", fgColor=SLATE)
ws_outputs["B8"] = "=Inputs!$B$2"
ws_outputs["B8"].number_format = '#,##0'
ws_outputs["B8"].font = Font(bold=True, color=NAVY)
ws_outputs["B8"].fill = PatternFill("solid", fgColor=LIGHT_BLUE)
ws_outputs["C8"] = "Edit Inputs!B2 to switch the report scenario."
ws_outputs["C8"].font = Font(italic=True, color=SLATE)
ws_outputs.merge_cells("C8:E8")

report_headers = ["Customer Name", "City", "Product", "Order Amount", "Order ID"]
for col_num, header in enumerate(report_headers, start=1):
    cell = ws_outputs.cell(10, col_num, header)
    cell.font = Font(bold=True, color=WHITE)
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.alignment = Alignment(horizontal="center")

for row_num in range(11, 16):
    position = f"ROWS($A$11:A{row_num})"
    formulas = [
        f'=IFERROR(INDEX(Calcs!$E$4:$E$8,MATCH({position},Calcs!$J$4:$J$8,0)),"")',
        f'=IFERROR(INDEX(Calcs!$F$4:$F$8,MATCH({position},Calcs!$J$4:$J$8,0)),"")',
        f'=IFERROR(INDEX(Calcs!$C$4:$C$8,MATCH({position},Calcs!$J$4:$J$8,0)),"")',
        f'=IFERROR(INDEX(Calcs!$D$4:$D$8,MATCH({position},Calcs!$J$4:$J$8,0)),"")',
        f'=IFERROR(INDEX(Calcs!$A$4:$A$8,MATCH({position},Calcs!$J$4:$J$8,0)),"")',
    ]
    for col_num, formula in enumerate(formulas, start=1):
        ws_outputs.cell(row_num, col_num, formula)
        ws_outputs.cell(row_num, col_num).border = Border(bottom=thin_gray)
    ws_outputs.cell(row_num, 4).number_format = '#,##0'

ws_outputs.freeze_panes = "A11"

for ws in (ws_inputs, ws_calcs, ws_outputs):
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None:
                cell.alignment = Alignment(vertical="center", wrap_text=cell.alignment.wrap_text)
    for col_letter, width in {"A":18, "B":18, "C":18, "D":18, "E":20, "F":16, "G":16, "H":18, "I":22, "J":16}.items():
        ws.column_dimensions[col_letter].width = width

# Make the input controls visually obvious.
ws_inputs["B2"].fill = PatternFill("solid", fgColor=TEAL)
ws_inputs["B2"].font = Font(bold=True, color=NAVY)
ws_inputs["B2"].alignment = Alignment(horizontal="center")
ws_inputs.freeze_panes = "A6"
ws_calcs.freeze_panes = "A4"

# Document workbook calculation behavior.
wb.calculation.fullCalcOnLoad = True
wb.calculation.forceFullCalc = True
wb.calculation.calcMode = "auto"

wb.save(OUTPUT)

# Verification: reopen and check workbook structure, formulas, and formula restrictions.
check = load_workbook(OUTPUT, data_only=False)
assert check.sheetnames == ["Inputs", "Calcs", "Outputs"]
formula_cells = []
for sheet in check.worksheets:
    for row in sheet.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith("="):
                formula_cells.append(cell.value)
                assert "." not in cell.value, f"Dot found in formula {cell.coordinate}: {cell.value}"
                assert "@" not in cell.value, f"At-sign found in formula {cell.coordinate}: {cell.value}"
assert any("Inputs!$B$2" in formula for formula in formula_cells)
assert any("XLOOKUP" in formula for formula in formula_cells)
assert any("INDEX" in formula and "MATCH" in formula for formula in formula_cells)
assert any("IFERROR" in formula for formula in formula_cells)
print(f"Created and verified {OUTPUT} with {len(formula_cells)} formulas.")
